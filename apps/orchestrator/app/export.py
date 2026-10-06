"""
Word exports in Phoneme's corporate document format (2026-10-06).

Follows the Phoneme Proposal skill's VERIFIED corporate template (matched
against Phoneme-issued documents such as the J&K High Court HLD/LLD and the
Teamora BRD/PRD):
  * cover page: centred Phoneme logo, orange banner (product name, tagline,
    document subtitle) with a grey bar beneath, document-type label, and near
    the bottom the project code / month-year / "Submitted by" / security
    classification -- no version table on the cover;
  * Document Control page: Document Control, Authorization, Change Log
    (Version | Date | Section | A/M/D | Description | Reviewed By), Client
    Sign Off and the Confidentiality Agreement in dark red;
  * Table of Contents built from bookmarks + PAGEREF fields (renders in Word
    and headless LibreOffice -- no "update fields" step needed);
  * numbered sections, H1 orange 15pt (no rule), H2 dark grey 12pt, H3 black;
  * header: Phoneme logo left, grey document-type label right (no rule);
    footer: "Confidential | <Product> | Page X of Y" (no rule);
  * thin dark-grey page border on every page, Calibri throughout.

File and folder names follow the convention already used in the HR
Management workspace: <Product>/Requirement/<Product>_BRD_PRD_v<ver>.docx,
<Product>/Technical/<Product>_TechDesign_v<ver>.docx, <Product>/UI-UX/...
"""
import io
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

COMPANY = "Phoneme Solutions Pvt. Ltd."
PREPARED_BY = "GiveWings SDLC Platform (AI-assisted)"
ORANGE, DARK_RED, GREY_H2, GREY_TXT, GREY_BAR, LABEL_BG = "D34817", "9B2D1F", "404040", "6E6E6E", "8C8C8C", "F2F2F2"
FONT = "Calibri"
LOGO = Path(__file__).resolve().parent / "assets" / "phoneme_logo.png"
CONFIDENTIALITY = (
    "This document is copyrighted, and all rights are reserved. This document may not, in whole or in part, "
    "be copied, photocopied, reproduced, translated, or reduced to any electronic medium or machine-readable "
    "form without prior consent, in writing. This document is for internal use only and may, in whole or in "
    "part, be provided to anyone outside of the Company, including customers, clients, or prospects after "
    "taking an approval from an authorized representative."
)

# Folder layout per document type (matches the existing Phoneme workspaces).
FOLDERS = {"brdprd": "Requirement", "techdesign": "Technical", "uiux": "UI-UX"}


def product_code(state) -> str:
    name = re.sub(r"[^A-Za-z]", "", state.selected_name or "PROD").upper()
    return name[:4] or "PROD"


def product_slug(state) -> str:
    """'RelayReel' -> 'RelayReel', 'Content Box' -> 'Content_Box' (file-safe)."""
    s = re.sub(r"[^A-Za-z0-9]+", "_", (state.selected_name or "Product").strip()).strip("_")
    return s or "Product"


def project_code(state) -> str:
    return f"PHN-{product_code(state)}-{datetime.now(timezone.utc).year}-01"


def file_name(state, doc_type: str, version: str) -> str:
    v = version if version and version[0].isdigit() else "draft"
    stem = {"brdprd": "BRD_PRD", "techdesign": "TechDesign", "uiux": "UI_UX_Mockups",
            "prototype": "Prototype", "stack": "Technical_Stack_Charter"}[doc_type]
    ext = {"uiux": "html", "prototype": "html", "stack": "md"}.get(doc_type, "docx")
    return f"{product_slug(state)}_{stem}_v{v}.{ext}"


def relative_path(state, doc_type: str, version: str) -> str:
    folder = FOLDERS.get(doc_type) or ("Technical" if doc_type == "stack" else "UI-UX")
    return f"{product_slug(state)}/{folder}/{file_name(state, doc_type, version)}"


def _date(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d-%b-%Y")
    except (TypeError, ValueError):
        return datetime.now(timezone.utc).strftime("%d-%b-%Y")


# ----------------------------------------------------------------- XML helpers
def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _table_borders(table, color: str | None = "BFBFBF", size: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if color:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _fix_widths(table, fractions, total_in: float = 6.47) -> None:
    """Fixed layout + tblGrid + per-cell widths: honoured by Word AND LibreOffice."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(total_in * 1440)))
    tw.set(qn("w:type"), "dxa")
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tw)
    grid = table._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for gc, f in zip(cols, fractions):
        gc.set(qn("w:w"), str(int(total_in * f * 1440)))
    for row in table.rows:
        for cell, f in zip(row.cells, fractions):
            cell.width = Inches(total_in * f)


def _page_border(section) -> None:
    sectPr = section._sectPr
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "page")
    pg.set(qn("w:display"), "allPages")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), GREY_H2)
        pg.append(el)
    # pgBorders must come after pgMar/pgSz in sectPr order; inserting before cols is valid.
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)


def _field(paragraph, instr: str, placeholder: str = "1", size=None, color=None, bold=False):
    """Complex field (PAGE, NUMPAGES, PAGEREF ...) with a cached result."""
    def run(child_tag=None, text=None, attrs=None):
        r = paragraph.add_run()
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r.bold = bold
        if child_tag:
            el = OxmlElement(child_tag)
            for k, v in (attrs or {}).items():
                el.set(qn(k), v)
            if text is not None:
                el.text = text
                el.set(qn("xml:space"), "preserve")
            r._r.append(el)
        return r
    run("w:fldChar", attrs={"w:fldCharType": "begin", "w:dirty": "true"})
    run("w:instrText", text=f" {instr} ")
    run("w:fldChar", attrs={"w:fldCharType": "separate"})
    r = paragraph.add_run(placeholder)
    if size:
        r.font.size = size
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    r.bold = bold
    run("w:fldChar", attrs={"w:fldCharType": "end"})


class PhonemeDoc:
    """Phoneme corporate document (proposal-skill template)."""

    def __init__(self, state, doc_type_label: str, title: str, subtitle: str):
        self.state = state
        self.label = doc_type_label          # e.g. "BRD / PRD", "HLD / LLD"
        self.title, self.subtitle = title, subtitle
        self.doc = Document()
        self.toc: list[tuple[int, str, str]] = []
        self._bm = 0
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.9))
        sec.top_margin, sec.bottom_margin = Inches(0.9), Inches(0.8)
        _page_border(sec)

    # ------------------------------------------------------------ styles
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(6)
        for lvl, size, color in ((1, 15, ORANGE), (2, 12, GREY_H2), (3, 11, "000000")):
            h = st[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
            h.font.italic = False
            h.font.color.rgb = RGBColor.from_string(color)
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(k), FONT)
            for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(k)) is not None:
                    del rf.attrib[qn(k)]
            pf = h.paragraph_format
            pf.space_before, pf.space_after = Pt(14 if lvl == 1 else 10), Pt(6 if lvl == 1 else 4)
            pf.keep_with_next = True

    # ------------------------------------------------------------ primitives
    def _bookmark(self, paragraph, text: str) -> str:
        name = f"_PhnToc{self._bm}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bm))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bm))
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        self._bm += 1
        return name

    def h(self, text: str, level: int = 1, toc: bool = True, color: str | None = None):
        para = self.doc.add_heading(text, level=level)
        if color:
            for r in para.runs:
                r.font.color.rgb = RGBColor.from_string(color)
        if toc and level <= 2:
            self.toc.append((level, text, self._bookmark(para, text)))
        return para

    def p(self, text: str = "", bold=False, italic=False, size=None, color=None, align=None, justify=True):
        para = self.doc.add_paragraph()
        run = para.add_run(text or "")
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if align is not None:
            para.alignment = align
        elif justify and len(text or "") > 90:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def bullets(self, items):
        for it in items:
            if str(it).strip():
                self.doc.add_paragraph(str(it), style="List Bullet")

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.first_line_indent = Inches(-0.25)
            r = para.add_run(f"{i}.\t")
            r.bold = True
            para.add_run(str(it))

    def table(self, header, rows, widths=None, label_col=False, font_size=9.5):
        """Grey-bordered table, header row (and optional label column) shaded F2F2F2."""
        cols = len(header) if header else len(rows[0])
        t = self.doc.add_table(rows=0, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(t)
        all_rows = ([header] if header else []) + rows
        for ri, row in enumerate(all_rows):
            cells = t.add_row().cells
            for ci, txt in enumerate(row):
                c = cells[ci]
                c.text = ""
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                run = para.add_run(str(txt if txt is not None else ""))
                run.font.size = Pt(font_size)
                is_head = header is not None and ri == 0
                if is_head or (label_col and ci == 0):
                    _shade(c, LABEL_BG)
                if is_head:
                    run.bold = True
        if widths:
            _fix_widths(t, widths)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ front matter
    def cover(self):
        s = self.state
        d = self.doc
        for _ in range(3):
            d.add_paragraph()
        if LOGO.exists():
            lp = d.add_paragraph()
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lp.add_run().add_picture(str(LOGO), width=Inches(1.9))
        for _ in range(2):
            d.add_paragraph()
        banner = d.add_table(rows=2, cols=1)
        banner.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_borders(banner, None)
        top, bar = banner.rows[0].cells[0], banner.rows[1].cells[0]
        _shade(top, ORANGE)
        _shade(bar, GREY_BAR)
        top.text = ""
        lines = [(self.title, 20, True, False)]
        if s.brand and s.brand.tagline:
            lines.append((s.brand.tagline, 10, False, True))
        if self.subtitle:
            lines.append((self.subtitle, 13, False, False))
        first = True
        for text, size, bold, italic in lines:
            para = top.paragraphs[0] if first else top.add_paragraph()
            first = False
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(10 if bold else 2)
            para.paragraph_format.space_after = Pt(2)
            r = para.add_run(text)
            r.bold, r.italic = bold, italic
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        top.add_paragraph().paragraph_format.space_after = Pt(4)
        bar.text = ""
        bar.paragraphs[0].paragraph_format.space_after = Pt(0)
        bar.paragraphs[0].add_run(" ").font.size = Pt(4)
        trPr = banner.rows[1]._tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), "140")
        h.set(qn("w:hRule"), "exact")
        trPr.append(h)
        d.add_paragraph()
        lab = d.add_paragraph()
        lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lr = lab.add_run(self.label)
        lr.bold, lr.font.size = True, Pt(14)
        for _ in range(11):
            d.add_paragraph()
        now = datetime.now(timezone.utc)
        for text, kw in (
            (f"PROJECT CODE: {project_code(s)}", dict(bold=True, color=ORANGE, size=10.5)),
            (now.strftime("%B %Y"), dict(size=10)),
            (f"Submitted by: {COMPANY}", dict(size=10)),
            ("Security Classification: Confidential — Internal", dict(italic=True, color=GREY_TXT, size=9)),
        ):
            para = self.p(text, align=WD_ALIGN_PARAGRAPH.CENTER, **kw)
            para.paragraph_format.space_after = Pt(2)

    def start_body(self):
        """New section for everything after the cover: header + footer + border."""
        d = self.doc
        sec = d.add_section(WD_SECTION.NEW_PAGE)
        _page_border(sec)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        # cover section keeps an empty header/footer
        first = d.sections[0]
        first.header.is_linked_to_previous = False
        first.footer.is_linked_to_previous = False
        # header: logo left, label right (borderless 2-cell table)
        hdr = sec.header
        hdr.paragraphs[0].text = ""
        ht = hdr.add_table(rows=1, cols=2, width=Inches(6.47))
        _table_borders(ht, None)
        lc, rc = ht.rows[0].cells
        if LOGO.exists():
            lc.paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(1.15))
        rp = rc.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rr = rp.add_run(self.label)
        rr.font.size, rr.font.color.rgb = Pt(8), RGBColor.from_string(GREY_TXT)
        # footer: Confidential | Product | Page X of Y
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size, col = Pt(8), GREY_TXT

        def txt(t):
            r = fp.add_run(t)
            r.font.size, r.font.color.rgb = size, RGBColor.from_string(col)
        txt(f"Confidential | {self.state.selected_name or 'Product'} | Page ")
        _field(fp, "PAGE", "1", size=size, color=col)
        txt(" of ")
        _field(fp, "NUMPAGES", "1", size=size, color=col)

    def document_control(self, description: str, identification: str, location: str, change_rows,
                         prepared_on: str, reviewed: str, approved: str):
        self.h("Document Control", 1, toc=False)
        self.table(["Field", "Value"], [
            ["Document Description", description],
            ["Identification", identification],
            ["Security Classification", "Confidential — Internal"],
            ["Location", location],
        ], widths=[0.3, 0.7], label_col=True)
        self.h("Authorization", 1, toc=False)
        self.table(["", "Name of the Person", "Date"], [
            ["Prepared by", PREPARED_BY, prepared_on],
            ["Reviewed by", reviewed, ""],
            ["Approved by", approved, ""],
        ], widths=[0.22, 0.5, 0.28], label_col=True)
        self.h("Change Log", 1, toc=False)
        self.table(["Version", "Date", "Section", "A/M/D", "Description", "Reviewed By"], change_rows,
                   widths=[0.09, 0.13, 0.12, 0.1, 0.41, 0.15], font_size=9)
        self.h("Client Sign Off", 1, toc=False)
        self.table(["Stakeholder Name", "Role", "Sign Off Date", "Communication"],
                   [["", "", "", ""], ["", "", "", ""]], widths=[0.3, 0.25, 0.2, 0.25])
        self.h("Confidentiality Agreement", 2, toc=False, color=DARK_RED)
        self.p(CONFIDENTIALITY, size=10)
        self.page_break()
        self._toc_anchor_heading = self.h("Table of Contents", 1, toc=False)
        self._toc_anchor = self.doc.add_paragraph()
        self.page_break()

    def finish_toc(self):
        """Insert TOC lines (bookmarks + PAGEREF) now that every heading exists."""
        anchor = self._toc_anchor
        for level, text, bm in self.toc:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(1 if level == 2 else 3)
            pf.left_indent = Inches(0.25 if level == 2 else 0)
            pf.tab_stops.add_tab_stop(Inches(6.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = para.add_run(text)
            r.bold = level == 1
            r.font.size = Pt(10 if level == 1 else 9.5)
            para.add_run("\t").font.size = Pt(9.5)
            _field(para, f"PAGEREF {bm} \\h", "", size=Pt(9.5 if level == 2 else 10), bold=level == 1)
            anchor._p.addprevious(para._p)

    def bytes(self) -> bytes:
        self.finish_toc()
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


# ----------------------------------------------------------------- shared bits
def _baselines(state, doc_type):
    return [b for b in state.baselines if b.doc_type == doc_type]


def _version(state, doc_type: str):
    b = _baselines(state, doc_type)
    return (b[-1].version, "Approved Baseline") if b else ("Draft", "Draft — pending Reviewed/Approved sign-off")


def _change_rows(state, doc_type, section_of=lambda ids: "All"):
    rows = []
    for i, b in enumerate(_baselines(state, doc_type)):
        changed = [x.strip() for x in b.description.split(":", 1)[1].split(",")] if i and ":" in b.description else []
        rows.append([b.version, _date(b.at), section_of(changed) if changed else "All", "M" if i else "A",
                     b.description, "Product owner (approved in GiveWings)"])
    if not rows:
        rows.append(["0.1", _date(datetime.now(timezone.utc).isoformat()), "All", "A",
                     "Working draft generated in GiveWings — not yet baselined", "Pending"])
    return rows


def _reviewers(state, doc_type):
    b = _baselines(state, doc_type)
    return ("Product owner — reviewed in GiveWings", f"Product owner — baselined v{b[-1].version} on {_date(b[-1].at)}") if b else ("Pending", "Pending")


# ----------------------------------------------------------------- BRD / PRD
def brdprd_docx(state, reqs) -> tuple[bytes, str]:
    version, status = _version(state, "brdprd")
    fname = file_name(state, "brdprd", version)
    code = project_code(state)
    brief = state.concept_brief
    specs = state.module_specs or []
    std = [m for m in specs if m.kind == "standard"]
    biz = [m for m in specs if m.kind != "standard"]
    b = PhonemeDoc(state, "BRD / PRD", state.selected_name or "Product",
                   "Business & Product Requirements")
    b.cover()
    b.start_body()
    order = {r.req_id: i for i, r in enumerate(reqs, 1)}
    reviewed, approved = _reviewers(state, "brdprd")
    b.document_control(
        f"Business Requirements Document / Product Requirements Document (BRD/PRD) — {state.selected_name}",
        f"{code}-BRDPRD, Version {version} — {status}",
        relative_path(state, "brdprd", version),
        _change_rows(state, "brdprd", lambda ids: ", ".join(f"4.{order[i]}" for i in ids if i in order) or "All"),
        _date(datetime.now(timezone.utc).isoformat()), reviewed, approved,
    )

    b.h("1. Introduction", 1)
    b.h("1.1 Overview", 2)
    b.p((brief.summary if brief and brief.summary else state.concept_summary) or "")
    if state.research and state.research.positioning_reframe:
        b.p(state.research.positioning_reframe, italic=True)
    b.h("1.2 Purpose", 2)
    b.p(f"This document records what {state.selected_name} must do, for whom and why, in business language. "
        "It carries no technical design, pricing or delivery plan — the Technical Design is produced from this "
        "document once it is baselined, and every module there traces back to a requirement here.")
    b.h("1.3 Scope", 2)
    b.p("In scope:", bold=True)
    b.bullets([f"{m.name} — {m.description}" if m.description else m.name for m in biz] or [r.module for r in reqs])
    if std:
        b.p("Standard GiveWings modules included in every product: " + ", ".join(m.name for m in std) + ".")
    out = list(brief.out_of_scope) if brief else []
    deferred = [f"{d.question} (deferred)" for r in reqs for d in r.decisions if d.deferred]
    if out or deferred:
        b.p("Out of scope / deferred:", bold=True)
        b.bullets(out + deferred)
    b.h("1.4 References", 2)
    b.bullets([f"Market research and concept brief captured in GiveWings for {state.selected_name}",
               "GiveWings standard module templates v1.0 (sign-in, profile & settings, admin & roles, security, audit & health)"]
              + (["Digital Personal Data Protection Act, 2023 (India)"] if brief and "india" in (brief.market or "").lower() else []))

    b.h("2. Overall Description / Background", 1)
    b.h("2.1 Product Context", 2)
    if brief:
        b.table(["Item", "Decision"], [[k, v] for k, v in (
            ("Target users", brief.target_users), ("Core workflow", brief.core_workflow),
            ("Input channels", brief.input_channels), ("Privacy & visibility", brief.privacy_mode),
            ("Launch market", brief.market), ("Launch platforms", brief.platforms),
        ) if v], widths=[0.3, 0.7], label_col=True)
    if state.research and state.research.killer_feature:
        b.h("2.2 Differentiator", 2)
        b.p(state.research.killer_feature)

    b.h("3. Stakeholders & Personas", 1)
    actors = []
    for r in reqs:
        for a in (r.doc.actors if r.doc else []):
            if a not in actors:
                actors.append(a)
    b.table(["Role", "Involved in"], [[a, ", ".join(r.req_id for r in reqs if r.doc and a in r.doc.actors)] for a in actors]
            or [["—", "—"]], widths=[0.35, 0.65], label_col=True)

    b.h("4. Specific Requirements", 1)
    access = {m.name: m for m in specs}
    for r in reqs:
        n = order[r.req_id]
        d = r.doc
        b.h(f"4.{n} {r.title} ({r.req_id})", 2)
        meta = [f"Module: {r.module}"]
        if r.standard_version:
            meta.append(f"GiveWings standard v{r.standard_version}")
        m = access.get(r.module)
        if m and m.kind != "standard":
            meta.append({"public": "Public — no sign-in", "mixed": "Public trial + sign-in"}.get(m.access, "Signed-in users"))
            if m.starts_when or m.outcome:
                meta.append(f"Starts when: {m.starts_when or '—'} → Outcome: {m.outcome or '—'}")
        meta.append(f"Status: {r.status}")
        b.p("  ·  ".join(meta), italic=True, size=9, color=GREY_TXT, justify=False)
        if not d:
            b.p(r.body)
            continue
        if d.summary:
            b.p(d.summary)
        if d.receives_from or d.hands_off_to:
            b.table(["Receives from", "Hands off to"], [[d.receives_from or "—", d.hands_off_to or "—"]], widths=[0.5, 0.5])
        if d.actors:
            b.p("Actors: " + ", ".join(d.actors), size=10)
        if d.journey:
            b.h("How it works", 3, toc=False)
            b.numbered([f"{s.title}{f' ({s.actor})' if s.actor else ''} — {s.detail}" for s in d.journey])
        if d.business_rules:
            b.h("Business rules", 3, toc=False)
            b.bullets(d.business_rules)
        if d.acceptance_criteria:
            b.h("Acceptance criteria", 3, toc=False)
            b.table(["#", "Criterion", "Priority"], [[a.id, a.criterion, a.priority] for a in d.acceptance_criteria],
                    widths=[0.08, 0.77, 0.15])
        if d.out_of_scope:
            b.h("Out of scope", 3, toc=False)
            b.bullets(d.out_of_scope)
        if r.decisions:
            b.h("Product-owner decisions", 3, toc=False)
            b.table(["Question", "Decision"], [[x.question, ("Deferred — " + x.answer if x.answer else "Deferred") if x.deferred else x.answer]
                                               for x in r.decisions], widths=[0.5, 0.5])
        if d.open_questions:
            b.h("Open questions", 3, toc=False)
            b.bullets(d.open_questions)

    market = (brief.market if brief else "") or "the launch market"
    b.h("5. Non-Functional Requirements", 1)
    b.table(["Category", "Requirement"], [
        ["Performance", "Everyday screens respond within 2 seconds for typical use; long-running AI work shows progress and never blocks the user."],
        ["Security", "Covered by the Security, Audit & Health standard module: protected data, immutable audit trail, role-based access."],
        ["Scalability", "The product grows with users and content volume without redesign; background work is queued, not done inline."],
        ["Availability", "Every module is watched by the health agent; outages alert the owner and admins."],
        ["Compliance", f"Personal data is handled under the data-protection law of {market} (e.g. India's DPDP Act 2023 where applicable)."],
        ["Localization", f"Language and formats suitable for {market}."],
    ], widths=[0.22, 0.78], label_col=True)

    b.h("6. Acceptance Criteria", 1)
    total_ac = sum(len(r.doc.acceptance_criteria) for r in reqs if r.doc)
    b.p(f"The product is accepted when all {total_ac} acceptance criteria in Section 4 marked Must are met and verified, "
        "and the non-functional requirements in Section 5 are demonstrated in the staging environment.")

    b.h("7. Assumptions & Dependencies", 1)
    b.bullets((list(brief.assumptions) if brief else []) + [
        "Each product ships its own standard modules (sign-in, profile, admin, security & health) and can be hosted independently of GiveWings.",
        "Technical choices (stack, APIs, data model) are made in the Technical Design stage from this baselined document.",
    ])

    b.h("8. Success Metrics / KPIs", 1)
    b.p("Targets are to be confirmed by the product owner before launch.", italic=True, color=DARK_RED, size=10)
    rows = [[f"Users reaching: {m.outcome}", "To be confirmed", "Product analytics"] for m in biz if m.outcome]
    rows += [["Monthly active users", "To be confirmed", "Product analytics"], ["Module availability", "To be confirmed", "Health agent reports"]]
    b.table(["Metric", "Target", "How measured"], rows, widths=[0.5, 0.2, 0.3])

    b.h("9. Road Ahead — SDLC Sequence", 1)
    b.numbered(["BRD/PRD baselined (this document)", "Technical Design (HLD/LLD) from the baselined BRD/PRD",
                "UI/UX screens and clickable prototype in the product's brand", "Build, test and release"])
    return b.bytes(), fname


# ----------------------------------------------------------------- Technical Design
STACK_LABELS = [
    ("product_type", "Product type"), ("frontend", "Frontend"), ("backend", "Backend"),
    ("data", "Data"), ("ai_ml", "AI / ML components"), ("hosting", "Hosting & infra"),
    ("integrations", "Third-party integrations"), ("devops", "DevOps & CI/CD"),
    ("security", "Security & compliance baseline"), ("conventions", "Coding conventions"),
]


def techdesign_docx(state, reqs) -> tuple[bytes, str]:
    """Phoneme Technical Design Document (HLD/LLD). Module numbering mirrors
    the BRD/PRD Specific Requirements 1:1 (BRD 4.n -> TDD 2.2.n and 3.n)."""
    version, status = _version(state, "techdesign")
    fname = file_name(state, "techdesign", version)
    code = project_code(state)
    brd_v = [x.version for x in _baselines(state, "brdprd")]
    b = PhonemeDoc(state, "HLD / LLD", state.selected_name or "Product", "Technical Design Document (HLD/LLD)")
    b.cover()
    b.start_body()
    hld = next((t for t in state.tech_designs if t.kind == "hld"), None)
    llds = [t for t in state.tech_designs if t.kind == "lld" and t.doc]
    order = {r.req_id: i for i, r in enumerate(reqs, 1)}
    llds.sort(key=lambda t: order.get(t.req_id, 999))
    by_td = {t.td_id: order.get(t.req_id) for t in llds}
    reviewed, approved = _reviewers(state, "techdesign")
    b.document_control(
        f"Technical Design Document (HLD/LLD) — {state.selected_name}",
        f"{code}-TD, Version {version} — {status}",
        relative_path(state, "techdesign", version),
        _change_rows(state, "techdesign", lambda ids: ", ".join(f"3.{by_td[i]}" if by_td.get(i) else "2.1" for i in ids) or "All"),
        _date(datetime.now(timezone.utc).isoformat()), reviewed, approved,
    )

    b.h("1. Introduction", 1)
    b.h("1.1 Overview", 2)
    b.p((state.concept_brief.summary if state.concept_brief and state.concept_brief.summary else state.concept_summary) or "")
    b.h("1.2 Purpose", 2)
    b.p(f"This document defines how {state.selected_name} is built. It implements BRD/PRD v{brd_v[-1] if brd_v else 'draft'} "
        f"({code}-BRDPRD); every module section is numbered to match its BRD/PRD requirement for traceability.")
    b.h("1.3 References", 2)
    b.bullets([relative_path(state, "brdprd", brd_v[-1] if brd_v else "draft"),
               relative_path(state, "stack", (state.stack.confirmed_at and "1.0") or "draft") if state.stack else "Technical Stack Charter"])

    b.h("2. High-Level Design", 1)
    b.h("2.1 System Architecture", 2)
    if state.stack:
        b.p("Technical Stack Charter" + (" (confirmed)" if state.stack.confirmed else " (not yet confirmed)"), bold=True)
        b.table(["Area", "Decision"], [[label, getattr(state.stack, k) or "—"] for k, label in STACK_LABELS],
                widths=[0.3, 0.7], label_col=True)
    if hld and hld.doc:
        b.p(hld.doc.overview)
        if hld.doc.components:
            b.table(["Component", "Responsibility"], [[c.name, c.responsibility] for c in hld.doc.components], widths=[0.3, 0.7])
        if hld.doc.sequence:
            b.p("Main end-to-end flow", bold=True)
            b.numbered(hld.doc.sequence)
    b.h("2.2 Module Breakdown", 2)
    for t in llds:
        n = order.get(t.req_id, "")
        b.h(f"2.2.{n} {t.module} ({t.req_id} → {t.td_id})", 3, toc=False)
        b.p(t.doc.overview)
        if t.doc.components:
            b.bullets([f"{c.name} — {c.responsibility}" for c in t.doc.components])
    b.h("2.3 Integration Points", 2)
    integ = (hld.doc.integrations if hld and hld.doc else []) + [i for t in llds for i in t.doc.integrations]
    b.bullets(integ or ["None beyond the stack charter."])

    b.h("3. Low-Level Design", 1)
    for t in llds:
        n = order.get(t.req_id, "")
        d = t.doc
        b.h(f"3.{n} {t.module} ({t.td_id})", 2)
        b.h(f"3.{n}.1 Data model", 3, toc=False)
        if d.data_model:
            for e in d.data_model:
                b.p(f"{e.name}" + (f" — {e.description}" if e.description else ""), bold=True, justify=False)
                if e.fields:
                    b.table(["Field", "Type", "Notes"], [[f.name, f.type, f.notes] for f in e.fields], widths=[0.3, 0.2, 0.5])
        else:
            b.p("No module-specific data.", italic=True)
        b.h(f"3.{n}.2 API contracts", 3, toc=False)
        if d.apis:
            b.table(["Method", "Endpoint", "Purpose", "Request", "Response"],
                    [[a.method, a.path, a.purpose, a.request, a.response] for a in d.apis],
                    widths=[0.09, 0.25, 0.24, 0.21, 0.21], font_size=8.5)
        b.h(f"3.{n}.3 Key sequence flow", 3, toc=False)
        b.numbered(d.sequence or ["—"])
        b.h(f"3.{n}.4 Edge cases & error handling", 3, toc=False)
        b.bullets(d.edge_cases or ["—"])
        if t.decisions:
            b.p("Decisions recorded during review", bold=True, justify=False)
            b.table(["Question", "Decision"], [[x.question, ("Deferred — " + x.answer if x.answer else "Deferred") if x.deferred else x.answer]
                                               for x in t.decisions], widths=[0.5, 0.5])

    b.h("4. Security & Compliance", 1)
    sec = (hld.doc.security if hld and hld.doc else []) + [f"{t.module}: {s}" for t in llds for s in t.doc.security]
    b.bullets(sec or ["See the Security, Audit & Health module."])

    b.h("5. Non-Functional Design", 1)
    nfr = (hld.doc.nfr if hld and hld.doc else []) + [x for t in llds for x in t.doc.nfr]
    b.table(["Requirement (BRD/PRD)", "How it is achieved"], [[x.requirement, x.approach] for x in nfr] or [["—", "—"]], widths=[0.4, 0.6])

    b.h("6. Key Risks & Open Questions", 1)
    risks = (hld.doc.risks if hld and hld.doc else []) + [f"{t.module}: {r}" for t in llds for r in t.doc.risks]
    oq = [f"{t.td_id}: {q}" for t in state.tech_designs if t.doc for q in t.doc.open_questions]
    b.table(["Item", "Owner", "Target resolution"], [[x, "Product owner / architect", "Before build start"] for x in risks + oq]
            or [["None open", "—", "—"]], widths=[0.6, 0.2, 0.2])

    b.h("7. Road Ahead", 1)
    b.numbered(["UI/UX screens and clickable prototype built on these API contracts",
                "Sprint planning per module (3.x) in BRD/PRD order", "Build, test against the BRD/PRD acceptance criteria, release"])
    return b.bytes(), fname


# ----------------------------------------------------------------- Stack charter (markdown)
def stack_charter_md(state) -> tuple[bytes, str]:
    s = state.stack
    v = "1.0" if s and s.confirmed else "draft"
    lines = [f"# {state.selected_name} — Technical Stack Charter v{v}", "",
             f"Project code: {project_code(state)}  ", f"Status: {'Confirmed ' + _date(s.confirmed_at) if s and s.confirmed else 'Draft'}  ",
             f"Submitted by: {COMPANY}", "", "| Area | Decision |", "|---|---|"]
    for k, label in STACK_LABELS:
        val = (getattr(s, k) if s else "") or "TBD"
        lines.append(f"| {label} | {val.replace('|', '/').replace(chr(10), ' ')} |")
    if s and s.changelog:
        lines += ["", "## Change log", ""] + [f"- {c}" for c in s.changelog]
    return "\n".join(lines).encode(), file_name(state, "stack", v)
